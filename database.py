import sqlite3
import hashlib
from datetime import datetime

from pathlib import Path

DB_FILE = str((Path(__file__).parent / "silambam_attendance.db").resolve())

def get_connection():
    """Returns a SQLite connection with dict-like row accessibility."""
    conn = sqlite3.connect(DB_FILE, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn

def hash_password(password, salt="kalaithai_silambam_salt_2026"):
    """Hash password using SHA256 with salt."""
    return hashlib.sha256((password + salt).encode()).hexdigest()

def init_db():
    """Initialize database tables and default accounts."""
    conn = get_connection()
    c = conn.cursor()

    # Users Table
    c.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            full_name TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'trainer',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    # Students Table
    c.execute("""
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            reg_no TEXT UNIQUE NOT NULL,
            name TEXT NOT NULL,
            category TEXT,
            gender TEXT,
            phone TEXT,
            batch TEXT DEFAULT 'Morning (06:30 AM - 08:30 AM)',
            status TEXT DEFAULT 'active',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    # Attendance Table
    c.execute("""
        CREATE TABLE IF NOT EXISTS attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER NOT NULL,
            date TEXT NOT NULL,
            status TEXT NOT NULL,
            batch TEXT DEFAULT 'Morning (06:30 AM - 08:30 AM)',
            marked_by TEXT NOT NULL,
            timestamp TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (student_id) REFERENCES students(id) ON DELETE CASCADE,
            UNIQUE(student_id, date)
        )
    """)

    # Migrations for existing database
    try:
        c.execute("ALTER TABLE students ADD COLUMN batch TEXT DEFAULT 'Morning (06:30 AM - 08:30 AM)'")
    except sqlite3.OperationalError:
        pass

    try:
        c.execute("ALTER TABLE attendance ADD COLUMN batch TEXT DEFAULT 'Morning (06:30 AM - 08:30 AM)'")
    except sqlite3.OperationalError:
        pass

    c.execute("UPDATE students SET batch = 'Morning (06:30 AM - 08:30 AM)' WHERE batch IS NULL OR batch = ''")
    c.execute("UPDATE attendance SET batch = 'Morning (06:30 AM - 08:30 AM)' WHERE batch IS NULL OR batch = ''")

    # Seed Default Users (Admin & Trainer)
    seed_default_users(c)

    conn.commit()
    conn.close()

def seed_default_users(cursor):
    """Seed initial Admin and Trainer accounts if not existing or ensure default passwords."""
    default_users = [
        ('admin', 'admin123', 'Master Admin', 'admin'),
        ('trainer', 'trainer123', 'Head Trainer', 'trainer')
    ]
    for username, password, full_name, role in default_users:
        pw_hash = hash_password(password)
        cursor.execute("SELECT id FROM users WHERE LOWER(TRIM(username)) = ?", (username.lower(),))
        row = cursor.fetchone()
        if row:
            uid = row[0] if isinstance(row, tuple) else row['id']
            cursor.execute(
                "UPDATE users SET password_hash = ?, full_name = ?, role = ? WHERE id = ?",
                (pw_hash, full_name, role, uid)
            )
        else:
            cursor.execute(
                "INSERT INTO users (username, password_hash, full_name, role) VALUES (?, ?, ?, ?)",
                (username, pw_hash, full_name, role)
            )

# --- USER MANAGEMENT FUNCTIONS ---

def add_user(username, password, full_name, role='trainer'):
    conn = get_connection()
    c = conn.cursor()
    try:
        c.execute(
            "INSERT INTO users (username, password_hash, full_name, role) VALUES (?, ?, ?, ?)",
            (username, hash_password(password), full_name, role)
        )
        conn.commit()
        conn.close()
        return True, "User registered successfully!"
    except sqlite3.IntegrityError:
        conn.close()
        return False, "Username already exists!"

def get_all_users():
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT id, username, full_name, role, created_at FROM users ORDER BY id DESC")
    rows = c.fetchall()
    conn.close()
    return [dict(r) for r in rows]

def delete_user(user_id):
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT username FROM users WHERE id = ?", (user_id,))
    user = c.fetchone()
    if user and user['username'] in ['admin', 'trainer']:
        conn.close()
        return False, "Protected system accounts cannot be deleted!"
    
    c.execute("DELETE FROM users WHERE id = ?", (user_id,))
    conn.commit()
    conn.close()
    return True, "User deleted successfully!"

# --- STUDENT MANAGEMENT FUNCTIONS ---

def add_student(reg_no, name, category, gender, phone, batch="Morning (06:30 AM - 08:30 AM)"):
    conn = get_connection()
    c = conn.cursor()
    try:
        c.execute(
            "INSERT INTO students (reg_no, name, category, gender, phone, batch) VALUES (?, ?, ?, ?, ?, ?)",
            (reg_no, name, category, gender, phone, batch)
        )
        conn.commit()
        conn.close()
        return True, "Student registered successfully!"
    except sqlite3.IntegrityError:
        conn.close()
        return False, "Registration Number already exists!"

def get_all_students(status='active'):
    conn = get_connection()
    c = conn.cursor()
    if status == 'all':
        c.execute("SELECT * FROM students ORDER BY name ASC")
    else:
        c.execute("SELECT * FROM students WHERE status = ? ORDER BY name ASC", (status,))
    rows = c.fetchall()
    conn.close()
    return [dict(r) for r in rows]

def delete_student(student_id):
    conn = get_connection()
    c = conn.cursor()
    c.execute("DELETE FROM students WHERE id = ?", (student_id,))
    c.execute("DELETE FROM attendance WHERE student_id = ?", (student_id,))
    conn.commit()
    conn.close()
    return True, "Student removed successfully!"

# --- ATTENDANCE FUNCTIONS ---

def mark_attendance(student_id, date_str, status, marked_by, batch="Morning (06:30 AM - 08:30 AM)"):
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        INSERT INTO attendance (student_id, date, status, marked_by, batch)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(student_id, date) DO UPDATE SET
            status = excluded.status,
            marked_by = excluded.marked_by,
            batch = excluded.batch,
            timestamp = CURRENT_TIMESTAMP
    """, (student_id, date_str, status, marked_by, batch))
    conn.commit()
    conn.close()

def get_attendance_by_date(date_str):
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        SELECT s.id, s.reg_no, s.name, s.category, 
               COALESCE(s.batch, 'Morning (06:30 AM - 08:30 AM)') as batch,
               COALESCE(a.status, 'Not Marked') as status,
               a.marked_by, a.timestamp
        FROM students s
        LEFT JOIN attendance a ON s.id = a.student_id AND a.date = ?
        WHERE s.status = 'active'
        ORDER BY s.name ASC
    """, (date_str,))
    rows = c.fetchall()
    conn.close()
    return [dict(r) for r in rows]

def get_individual_report(student_id):
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        SELECT date, status, COALESCE(batch, 'Morning (06:30 AM - 08:30 AM)') as batch, marked_by, timestamp
        FROM attendance
        WHERE student_id = ?
        ORDER BY date DESC
    """, (student_id,))
    rows = c.fetchall()
    conn.close()
def get_monthly_attendance_matrix(year_month):
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        SELECT s.reg_no, s.name, a.date, a.status
        FROM students s
        JOIN attendance a ON s.id = a.student_id
        WHERE a.date LIKE ? AND s.status = 'active'
        ORDER BY s.name, a.date
    """, (f"{year_month}%",))
    rows = c.fetchall()
    conn.close()
    return [dict(r) for r in rows]

def clear_all_students():
    """Wipe all student and attendance records so user can start fresh."""
    conn = get_connection()
    c = conn.cursor()
    c.execute("DELETE FROM attendance")
    c.execute("DELETE FROM students")
    conn.commit()
    conn.close()
    return True, "All students and attendance records cleared successfully."

import calendar
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def get_monthly_register_grid(year_month: str, batch: str = None):
    """
    Constructs a monthly attendance register grid for any given YYYY-MM.
    Columns: Reg No, Student Name, Batch, Day 1 .. Day N, Present Days, Absent Days, Rate %
    """
    try:
        year, month = map(int, year_month.split("-"))
    except Exception:
        return pd.DataFrame(), []

    num_days = calendar.monthrange(year, month)[1]

    conn = get_connection()
    c = conn.cursor()

    query = "SELECT id, reg_no, name, batch FROM students WHERE status = 'active'"
    params = []
    if batch and batch != "All":
        query += " AND batch = ?"
        params.append(batch)
    query += " ORDER BY name ASC"
    c.execute(query, params)
    students = [dict(r) for r in c.fetchall()]

    if not students:
        conn.close()
        return pd.DataFrame(), []

    # Get attendance for month
    c.execute("""
        SELECT student_id, date, status 
        FROM attendance 
        WHERE date LIKE ?
    """, (f"{year_month}%",))
    att_rows = c.fetchall()
    conn.close()

    att_map = {}
    for r in att_rows:
        att_map[(r["student_id"], r["date"])] = r["status"]

    day_cols = [f"{d:02d}" for d in range(1, num_days + 1)]
    records = []

    for s in students:
        s_id = s["id"]
        row = {
            "Reg No": s["reg_no"],
            "Student Name": s["name"],
            "Batch": s.get("batch", "Morning (06:30 AM - 08:30 AM)")
        }
        p_count = 0
        a_count = 0

        for day in range(1, num_days + 1):
            date_key = f"{year:04d}-{month:02d}-{day:02d}"
            status = att_map.get((s_id, date_key))
            day_str = f"{day:02d}"
            if status == "Present":
                row[day_str] = "P"
                p_count += 1
            elif status == "Absent":
                row[day_str] = "A"
                a_count += 1
            else:
                row[day_str] = "-"

        total_marked = p_count + a_count
        rate = round((p_count / total_marked * 100), 1) if total_marked > 0 else 0.0

        row["Present"] = p_count
        row["Absent"] = a_count
        row["Rate %"] = f"{rate}%"
        records.append(row)

    df = pd.DataFrame(records)
    return df, day_cols

def generate_monthly_excel(df: pd.DataFrame, year_month: str, batch_name: str = "All Batches") -> bytes:
    """Generates an Excel workbook for the monthly attendance register."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Register"

    header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFD700")
    title_font = Font(name="Calibri", size=13, bold=True, color="1E1B4B")
    sub_font = Font(name="Calibri", size=10, italic=True, color="4B5563")

    present_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    present_font = Font(name="Calibri", size=10, bold=True, color="065F46")

    absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    absent_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws.merge_cells("A1:Y1")
    ws["A1"] = "கலைத்தாய் சிலம்பம், குரும்பட்டி - KALAITHAI SILAMBAM, KURUMPATTI"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:Y2")
    ws["A2"] = f"MONTHLY ATTENDANCE REGISTER — Month: {year_month} | Batch: {batch_name}"
    ws["A2"].font = sub_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.append([])

    headers = list(df.columns)
    ws.append(headers)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, row_data in enumerate(df.values, 5):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            header_name = headers[col_idx - 1]

            if header_name in ["Reg No", "Student Name", "Batch"]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif val == "P":
                cell.fill = present_fill
                cell.font = present_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif val == "A":
                cell.fill = absent_fill
                cell.font = absent_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    from openpyxl.utils import get_column_letter
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        header_name = headers[col_idx - 1]
        if header_name == "Student Name":
            ws.column_dimensions[col_letter].width = 24
        elif header_name in ["Batch", "Timing"]:
            ws.column_dimensions[col_letter].width = 22
        elif header_name == "Reg No":
            ws.column_dimensions[col_letter].width = 10
        elif header_name in ["Present", "Absent", "Rate %"]:
            ws.column_dimensions[col_letter].width = 11
        else:
            ws.column_dimensions[col_letter].width = 5

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def bulk_import_students(df: pd.DataFrame, default_batch: str = "Morning (06:30 AM - 08:30 AM)"):
    """
    Bulk import students from uploaded pandas DataFrame (Excel or CSV).
    Maps columns: Reg No, Name, Category, Gender, Phone, Batch.
    """
    if df is None or df.empty:
        return False, 0, "Uploaded file is empty."

    col_map = {}
    for col in df.columns:
        norm = str(col).strip().lower().replace("_", " ").replace("-", " ")
        if norm in ["reg no", "regno", "roll no", "rollno", "id", "student id", "பதிவு எண்"]:
            col_map["reg_no"] = col
        elif norm in ["name", "student name", "studentname", "full name", "பெயர்"]:
            col_map["name"] = col
        elif norm in ["category", "cat", "பிரிவு"]:
            col_map["category"] = col
        elif norm in ["gender", "sex", "பாலினம்"]:
            col_map["gender"] = col
        elif norm in ["phone", "mobile", "contact", "phone number", "தொலைபேசி"]:
            col_map["phone"] = col
        elif norm in ["batch", "timing", "நேரம்"]:
            col_map["batch"] = col

    if "name" not in col_map:
        return False, 0, "Could not find 'Name' or 'Student Name' column in uploaded file."

    conn = get_connection()
    c = conn.cursor()

    c.execute("SELECT reg_no FROM students")
    existing_reg_nos = [r[0] for r in c.fetchall()]
    max_counter = 1
    for rno in existing_reg_nos:
        digits = "".join([ch for ch in str(rno) if ch.isdigit()])
        if digits:
            max_counter = max(max_counter, int(digits) + 1)

    imported_count = 0
    errors = []

    for idx, row in df.iterrows():
        name_val = row[col_map["name"]]
        if pd.isna(name_val) or not str(name_val).strip():
            continue
        name = str(name_val).strip()

        if "reg_no" in col_map and not pd.isna(row[col_map["reg_no"]]) and str(row[col_map["reg_no"]]).strip():
            reg_no = str(row[col_map["reg_no"]]).strip()
            if reg_no.endswith(".0"):
                reg_no = reg_no[:-2]
        else:
            reg_no = str(max_counter)
            max_counter += 1

        category = "Junior"
        if "category" in col_map and not pd.isna(row[col_map["category"]]):
            cat_val = str(row[col_map["category"]]).strip()
            if cat_val in ["Sub-Junior", "Junior", "Senior", "Super Senior"]:
                category = cat_val
            elif "sub" in cat_val.lower():
                category = "Sub-Junior"
            elif "super" in cat_val.lower():
                category = "Super Senior"
            elif "sen" in cat_val.lower():
                category = "Senior"
            else:
                category = cat_val

        gender = "Male"
        if "gender" in col_map and not pd.isna(row[col_map["gender"]]):
            gen_val = str(row[col_map["gender"]]).strip().capitalize()
            if gen_val in ["Male", "Female", "Other"]:
                gender = gen_val

        phone = ""
        if "phone" in col_map and not pd.isna(row[col_map["phone"]]):
            phone = str(row[col_map["phone"]]).strip()
            if phone.endswith(".0"):
                phone = phone[:-2]

        batch = default_batch
        if "batch" in col_map and not pd.isna(row[col_map["batch"]]):
            batch_val = str(row[col_map["batch"]]).strip()
            if batch_val:
                batch = batch_val

        try:
            c.execute("""
                INSERT INTO students (reg_no, name, category, gender, phone, batch, status)
                VALUES (?, ?, ?, ?, ?, ?, 'active')
                ON CONFLICT(reg_no) DO UPDATE SET
                    name = excluded.name,
                    category = excluded.category,
                    gender = excluded.gender,
                    phone = excluded.phone,
                    batch = excluded.batch,
                    status = 'active'
            """, (reg_no, name, category, gender, phone, batch))
            imported_count += 1
        except Exception as e:
            errors.append(f"Row {idx+1} ({name}): {str(e)}")

    conn.commit()
    conn.close()

    msg = f"Successfully imported {imported_count} students!"
    if errors:
        msg += f" (Errors on {len(errors)} rows: {'; '.join(errors[:2])})"
    return True, imported_count, msg

def get_student_sample_template():
    """Generates sample DataFrame for Excel/CSV template download."""
    data = [
        {"Reg No": "1", "Student Name": "Arun Kumar", "Category": "Junior", "Gender": "Male", "Phone": "9876543210", "Batch": "Morning (06:30 AM - 08:30 AM)"},
        {"Reg No": "2", "Student Name": "Kaviya Sri", "Category": "Sub-Junior", "Gender": "Female", "Phone": "9876543211", "Batch": "Morning (06:30 AM - 08:30 AM)"}
    ]
    return pd.DataFrame(data)